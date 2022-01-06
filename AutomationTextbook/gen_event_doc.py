import openpyxl as excel, docx, os

in_file = 'meibo.xlsx'
template_file = 'template-event.docx'
save_dir = os.path.join(
    os.path.dirname(__file__),
    'events')

if not os.path.exists(save_dir):
    os.mkdir(save_dir)

def read_book():
    result = []
    sheet = excel.load_workbook(in_file).active
    for row in sheet.iter_rows(min_row=2):
        v = [c.value for c in row]
        if v[0] is None: break
        result.append(v)
    return result

for person in read_book():
    name, zipno, addr = person
    card = {
        '住所:▲▲▲': addr,
        '●●様へ': name +'様へ'
    }
    doc = docx.Document(template_file)
    for p in doc.paragraphs:
        for k,v in card.items():
            if p.text.find(k) >= 0:
                p.text = p.text.replace(k, v)
                p.runs[0].font.bold = True
        save_file = os.path.join(save_dir,
                                 name+'樣.docx')
        print('save:', save_file)
        doc.save(save_file)